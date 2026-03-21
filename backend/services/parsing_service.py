from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from backend.schemas.extraction import RemittanceExtractionResult, RemittanceValidationReport
from backend.services.llm_service import extract_remittance_with_llm
from backend.utils.pdf import extract_pdf_text

BANK_COLS = {
    'Buchungsdatum': 'booking_date',
    'Valutadatum': 'value_date',
    'Betrag': 'amount',
    'Auftraggeber/Empfaenger': 'counterparty_name',
    'Auftraggeber/Empfanger': 'counterparty_name',
    'Auftraggeber/Empf\u00e4nger': 'counterparty_name',
    'Verwendungszweck': 'payment_purpose',
    'Auftraggeber-/Empfaengerkonto': 'counterparty_account',
    'Auftraggeber-/Empfangerkonto': 'counterparty_account',
    'Auftraggeber-/Empf\u00e4ngerkonto': 'counterparty_account',
    'BIC': 'bic',
    'Bank': 'bank_name',
    'Bankleitzahl': 'bank_code',
    'Buchungstext': 'booking_text',
    'Referenz': 'bank_reference',
    'IBAN': 'iban',
    'Konto': 'account_label',
    'Kontonummer': 'account_number',
    'Kundenreferenz': 'customer_reference',
    'Waehrung': 'currency',
    'Wahrung': 'currency',
    'W\u00e4hrung': 'currency',
}


def _as_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, date):
        return value
    return None


def parse_bank_statement(content: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    headers = [sheet.cell(1, i).value for i in range(1, sheet.max_column + 1)]

    mapped_indexes: dict[int, str] = {}
    for i, header in enumerate(headers, start=1):
        if header in BANK_COLS:
            mapped_indexes[i] = BANK_COLS[header]

    rows: list[dict] = []
    for row_idx in range(2, sheet.max_row + 1):
        row_data: dict[str, object] = {}
        for col_idx, target in mapped_indexes.items():
            row_data[target] = sheet.cell(row_idx, col_idx).value
        if not row_data or row_data.get('amount') is None:
            continue
        rows.append(
            {
                'line_number': row_idx - 1,
                'booking_date': _as_date(row_data.get('booking_date')),
                'value_date': _as_date(row_data.get('value_date')),
                'amount': Decimal(str(row_data.get('amount'))),
                'counterparty_name': row_data.get('counterparty_name'),
                'payment_purpose': row_data.get('payment_purpose'),
                'bank_reference': str(row_data.get('bank_reference')) if row_data.get('bank_reference') else None,
                'customer_reference': str(row_data.get('customer_reference')) if row_data.get('customer_reference') else None,
                'currency': str(row_data.get('currency') or 'EUR'),
            }
        )
    return rows


def validate_remittance_result(result: RemittanceExtractionResult) -> RemittanceValidationReport:
    reasons: list[str] = []

    if not result.lines:
        reasons.append('NO_LINE_ITEMS')

    for line in result.lines:
        if not line.raw_line_text:
            reasons.append('LINE_WITHOUT_RAW_TEXT')
            break

    line_total = sum((line.paid_amount or Decimal('0.00')) for line in result.lines)
    if result.total_paid_amount is not None and result.lines:
        if abs(result.total_paid_amount - line_total) > Decimal('0.02'):
            reasons.append('HEADER_LINE_SUM_MISMATCH')

    normalized = result.model_copy(deep=True)
    for line in normalized.lines:
        if line.currency is None:
            line.currency = normalized.document_currency or 'EUR'

    return RemittanceValidationReport(
        is_valid=not reasons,
        reason_codes=sorted(set(reasons)),
        normalized=normalized,
    )


def parse_remittance(content: bytes) -> tuple[RemittanceValidationReport, dict, str, str]:
    raw_text = extract_pdf_text(content)
    extracted, raw_llm_output, method, llm_model = extract_remittance_with_llm(raw_text, pdf_bytes=content)
    extracted.raw_text = raw_text
    report = validate_remittance_result(extracted)
    return report, raw_llm_output, method, llm_model
